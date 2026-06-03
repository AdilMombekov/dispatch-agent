"""Command handlers — each returns {"success": bool, "data": ..., "error": str|None}."""
import base64
import io
import json
import mimetypes
import os
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
from pathlib import Path

import psutil
from PIL import Image


def _ok(data) -> dict:
    return {"success": True, "data": data, "error": None}


def _err(msg: str) -> dict:
    return {"success": False, "data": None, "error": str(msg)}


# ── Screenshot ──────────────────────────────────────────────────────────────

def _capture_resized(max_size: int = 1280):
    """Grab primary monitor, resize to max_size, return (b64_jpeg, img_w, img_h, real_w, real_h)."""
    import mss

    with mss.mss() as sct:
        monitor = sct.monitors[1]  # primary monitor
        raw = sct.grab(monitor)
        img = Image.frombytes("RGB", raw.size, raw.bgra, "raw", "BGRX")

    real_w, real_h = img.width, img.height
    if img.width > max_size or img.height > max_size:
        img.thumbnail((max_size, max_size), Image.LANCZOS)

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=75)
    b64 = base64.b64encode(buf.getvalue()).decode()
    return b64, img.width, img.height, real_w, real_h


def handle_screenshot(_payload: dict) -> dict:
    try:
        b64, w, h, _, _ = _capture_resized()
        return _ok({"image": b64, "width": w, "height": h})
    except Exception as e:
        return _err(f"screenshot failed: {e}")


# ── Computer control (mouse / keyboard) — computer-use level ─────────────────

def _normalize_coords(payload: dict) -> None:
    """Mutate payload so payload['x'] and payload['y'] are plain numbers.

    Haiku-class routers sometimes pack coords as a single string ("920, 10"),
    a list ([920,10]), or under names like `coordinate`/`coordinates`. Accept
    those rather than erroring — the schema still says x/y separately, but a
    forgiving handler beats a frustrating chat loop."""
    def _split_pair(v):
        if isinstance(v, (list, tuple)) and len(v) >= 2:
            return v[0], v[1]
        if isinstance(v, str) and "," in v:
            parts = [p.strip() for p in v.split(",", 1)]
            if len(parts) == 2:
                return parts[0], parts[1]
        return None

    # Case 1: x or y itself is the pair (e.g. {"x": "920, 10"}).
    for key in ("x", "y"):
        pair = _split_pair(payload.get(key))
        if pair is not None:
            payload["x"], payload["y"] = pair
            return

    # Case 2: a sibling key carries the pair.
    for alt in ("coordinate", "coordinates", "coord", "coords", "xy", "position"):
        pair = _split_pair(payload.get(alt))
        if pair is not None:
            payload["x"], payload["y"] = pair
            return


def handle_computer(payload: dict) -> dict:
    """Control mouse/keyboard and return a fresh screenshot so the model can see.

    Coordinates (x, y) are interpreted in the SCREENSHOT pixel space (downscaled
    to max 1280) and scaled up to the real screen, so the model can click on what
    it sees in the image it received.
    """
    try:
        import pyautogui
        pyautogui.FAILSAFE = False

        action = payload.get("action", "")
        _normalize_coords(payload)
        real_w, real_h = pyautogui.size()
        longest = max(real_w, real_h)
        scale = longest / min(longest, 1280)  # image-space -> real-screen factor

        def to_real(x, y):
            return int(round(float(x) * scale)), int(round(float(y) * scale))

        if action == "screenshot":
            pass
        elif action in ("left_click", "right_click", "double_click", "move"):
            if payload.get("x") is None or payload.get("y") is None:
                return _err("x and y required for this action")
            rx, ry = to_real(payload["x"], payload["y"])
            pyautogui.moveTo(rx, ry, duration=0.1)
            if action == "left_click":
                pyautogui.click(rx, ry)
            elif action == "right_click":
                pyautogui.click(rx, ry, button="right")
            elif action == "double_click":
                pyautogui.doubleClick(rx, ry)
        elif action == "type":
            pyautogui.write(payload.get("text", ""), interval=0.02)
        elif action == "key":
            keys = (payload.get("keys") or payload.get("key") or "").strip()
            if not keys:
                return _err("keys required for action=key")
            if "+" in keys:
                pyautogui.hotkey(*[k.strip() for k in keys.split("+")])
            else:
                pyautogui.press(keys)
        elif action == "scroll":
            pyautogui.scroll(int(payload.get("amount", -400)))
        else:
            return _err(f"unknown computer action: {action}")

        time.sleep(0.5)  # let the UI settle before capturing
        b64, w, h, rw, rh = _capture_resized()
        return _ok({"action": action, "image": b64, "width": w, "height": h,
                    "screen_width": rw, "screen_height": rh})
    except Exception as e:
        return _err(f"computer error: {e}")


# ── Terminal ─────────────────────────────────────────────────────────────────

def handle_terminal(payload: dict) -> dict:
    # Accept both 'cmd' (new) and 'command' (legacy) keys
    cmd = payload.get("cmd") or payload.get("command", "")
    if not cmd:
        return _err("no command provided")
    # SAFETY.1: refuse hard-delete / wipe shell commands. Files must go through
    # safe_delete (trash backup), not be destroyed via the terminal.
    try:
        from agent.orchestrator.safety import detect_destructive
        hit = detect_destructive(cmd)
    except Exception:
        hit = None
    if hit:
        return _err(
            f"⛔ Заблокировано: команда содержит деструктивное удаление ('{hit}'). "
            "Удаление файлов только через мягкое удаление (/rm) — переносит в trash/ "
            "с возможностью восстановления.")
    try:
        result = subprocess.run(
            cmd,
            shell=True,
            capture_output=True,
            text=True,
            timeout=300,  # 5 min — protects against interactive prompts
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
            # cmd.exe writes in the console OEM codepage (e.g. cp866), not UTF-8
            encoding="oem" if sys.platform == "win32" else "utf-8",
            errors="replace",
        )
        output = result.stdout + result.stderr
        if len(output) > 3000:
            file_b64 = base64.b64encode(output.encode("utf-8")).decode()
            preview = output[:300].rstrip()
            return _ok({
                "output": preview + f"\n\n… ({len(output)} chars — full output attached)",
                "output_file_b64": file_b64,
                "filename": "terminal_output.txt",
                "returncode": result.returncode,
                "send_as_file": True,
            })
        return _ok({"output": output, "returncode": result.returncode, "send_as_file": False})
    except subprocess.TimeoutExpired:
        return _err("TIMEOUT: команда выполнялась дольше 5 минут и была убита. Возможно, процесс ждал ввода.")
    except Exception as e:
        return _err(f"terminal error: {e}")


# ── Launch App ────────────────────────────────────────────────────────────────

def handle_launch_app(payload: dict, apps: dict) -> dict:
    app_name = payload.get("app", "")
    app_path = apps.get(app_name, "")
    if not app_path:
        return _err(f"app '{app_name}' not configured or not found")
    try:
        subprocess.Popen(
            [app_path],
            creationflags=subprocess.DETACHED_PROCESS if sys.platform == "win32" else 0,
        )
        return _ok(f"Launched {app_name}")
    except Exception as e:
        return _err(f"could not launch {app_name}: {e}")


# ── Chrome Profile ────────────────────────────────────────────────────────────

def handle_chrome_profile(payload: dict, apps: dict) -> dict:
    profile_dir = payload.get("profile_directory", "")
    chrome_path = apps.get("chrome", "")
    if not chrome_path:
        return _err("Chrome not found in config")
    if not profile_dir:
        return _err("no profile_directory provided")
    try:
        subprocess.Popen(
            [chrome_path, f"--profile-directory={profile_dir}"],
            creationflags=subprocess.DETACHED_PROCESS if sys.platform == "win32" else 0,
        )
        return _ok(f"Launched Chrome with profile: {profile_dir}")
    except Exception as e:
        return _err(f"could not launch Chrome profile: {e}")


def handle_chrome_profiles_list(_payload: dict) -> dict:
    from agent.config import get_chrome_profiles
    return _ok(get_chrome_profiles())


# ── Install ────────────────────────────────────────────────────────────────────

def handle_install(payload: dict) -> dict:
    url = payload.get("url", "")
    if not url:
        return _err("no url provided")

    suffix = Path(url.split("?")[0]).suffix.lower()
    if suffix not in (".exe", ".msi", ".zip"):
        return _err(f"unsupported installer type: {suffix}")

    tmp_dir = Path(tempfile.mkdtemp())
    dest = tmp_dir / f"installer{suffix}"

    try:
        urllib.request.urlretrieve(url, dest)
    except Exception as e:
        return _err(f"download failed: {e}")

    try:
        if suffix == ".exe":
            subprocess.run(
                [str(dest), "/S", "/silent", "/quiet"],
                check=False,
                creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
            )
        elif suffix == ".msi":
            subprocess.run(
                ["msiexec", "/i", str(dest), "/quiet", "/norestart"],
                check=False,
                creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
            )
        elif suffix == ".zip":
            import zipfile
            extract_to = tmp_dir / "extracted"
            with zipfile.ZipFile(dest) as z:
                z.extractall(extract_to)
            return _ok(f"Extracted to {extract_to}")
        return _ok(f"Install initiated for {dest.name}")
    except Exception as e:
        return _err(f"install error: {e}")


# ── System Info ───────────────────────────────────────────────────────────────

def handle_system_info(_payload: dict) -> dict:
    try:
        cpu = psutil.cpu_percent(interval=0.5)
        ram = psutil.virtual_memory()
        disk = psutil.disk_usage("C:\\")
        return _ok({
            "cpu_percent": cpu,
            "ram_used_gb": round(ram.used / 1e9, 2),
            "ram_total_gb": round(ram.total / 1e9, 2),
            "ram_percent": ram.percent,
            "disk_used_gb": round(disk.used / 1e9, 2),
            "disk_total_gb": round(disk.total / 1e9, 2),
            "disk_percent": disk.percent,
        })
    except Exception as e:
        return _err(f"system-info error: {e}")


# ── Obsidian Context (direct FS read, multi-file + query filter) ─────────────

def handle_obsidian_context(payload: dict, obsidian_cfg: dict) -> dict:
    """Read one or more notes directly from the vault filesystem.

    payload keys:
      note   – single note name (string), OR
      notes  – list of note names (list[str])
      lines  – max tail lines per note (default 50)
      query  – optional keyword filter (case-insensitive grep)
    """
    notes_raw = payload.get("notes") or (
        [payload["note"]] if payload.get("note") else []
    )
    lines_limit = int(payload.get("lines", 50))
    query = payload.get("query", "").strip()

    vault_path = Path(obsidian_cfg.get("vault_path", "E:/Obsidian"))

    if not notes_raw:
        return _err("no note(s) specified")

    results: dict = {}
    for note_name in notes_raw:
        if not note_name.endswith(".md"):
            note_name += ".md"
        file_path = vault_path / note_name
        if not file_path.exists():
            results[note_name] = "Файл не найден."
            continue
        try:
            all_lines = file_path.read_text(encoding="utf-8").splitlines()
            if query:
                all_lines = [l for l in all_lines if query.lower() in l.lower()]
            tail = "\n".join(all_lines[-lines_limit:])
            results[note_name] = tail if tail.strip() else "Нет данных по запросу."
        except Exception as e:
            results[note_name] = f"Ошибка чтения: {e}"

    return _ok({"notes": results})


# ── Obsidian Log ──────────────────────────────────────────────────────────────

def _obsidian_ctx():
    """SSL context that accepts Obsidian's self-signed certificate."""
    import ssl
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    return ctx


def handle_obsidian_log(payload: dict, obsidian_cfg: dict) -> dict:
    """Append a log entry to an Obsidian note via the Local REST API plugin.

    payload keys:
      text  – required, the line to append
      note  – optional, override note name (default from config)
    """
    import urllib.request as _req
    import urllib.error as _uerr
    from datetime import datetime

    text = payload.get("text", "").strip()
    if not text:
        return _err("no text provided")

    host  = obsidian_cfg.get("host", "https://localhost:27124")
    token = obsidian_cfg.get("token", "")
    note  = payload.get("note") or obsidian_cfg.get("note", "claude97")
    vault = obsidian_cfg.get("vault", "")
    ctx   = _obsidian_ctx()

    ts   = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"- `{ts}` {text}\n"

    vault_path = f"/vault/{vault}/" if vault else "/vault/"
    url = f"{host}{vault_path}{note}.md"

    def _put(body: bytes):
        r = _req.Request(url, data=body, method="PUT",
                         headers={"Authorization": f"Bearer {token}",
                                  "Content-Type": "text/markdown"})
        with _req.urlopen(r, timeout=5, context=ctx) as resp:
            resp.read()

    def _patch(body: bytes):
        # Obsidian Local REST API v2: requires Target-Type + Heading + Operation
        r = _req.Request(url, data=body, method="PATCH",
                         headers={"Authorization": f"Bearer {token}",
                                  "Content-Type": "text/markdown",
                                  "Target-Type": "heading",
                                  "Heading": note,
                                  "Operation": "append"})
        with _req.urlopen(r, timeout=5, context=ctx) as resp:
            resp.read()

    try:
        _patch(line.encode("utf-8"))
        return _ok(f"Logged to Obsidian '{note}': {text[:80]}")
    except _uerr.HTTPError as e:
        if e.code == 404:
            try:
                _put((f"# {note}\n\n" + line).encode("utf-8"))
                return _ok(f"Created & logged to Obsidian '{note}'")
            except Exception as e2:
                return _err(f"obsidian create failed: {e2}")
        return _err(f"obsidian http {e.code}: {e}")
    except Exception as e:
        return _err(f"obsidian error: {e}")


def handle_obsidian_read(payload: dict, obsidian_cfg: dict) -> dict:
    """Read the last N lines of an Obsidian note."""
    import urllib.request as _req

    note  = payload.get("note") or obsidian_cfg.get("note", "claude97")
    lines = int(payload.get("lines", 20))
    host  = obsidian_cfg.get("host", "https://localhost:27124")
    token = obsidian_cfg.get("token", "")
    vault = obsidian_cfg.get("vault", "")
    ctx   = _obsidian_ctx()

    vault_path = f"/vault/{vault}/" if vault else "/vault/"
    url = f"{host}{vault_path}{note}.md"

    try:
        r = _req.Request(url, headers={"Authorization": f"Bearer {token}"})
        with _req.urlopen(r, timeout=5, context=ctx) as resp:
            content = resp.read().decode("utf-8")
        tail = "\n".join(content.splitlines()[-lines:])
        return _ok({"note": note, "content": tail})
    except Exception as e:
        return _err(f"obsidian read error: {e}")


# ── Send File ─────────────────────────────────────────────────────────────────

def handle_send_file(payload: dict) -> dict:
    # Accept several aliases the router model sometimes invents instead of `filepath`.
    filepath = (payload.get("filepath") or payload.get("path")
                or payload.get("file") or payload.get("file_path") or "").strip()
    if not filepath:
        return _err("no filepath provided")
    path = Path(filepath)
    if not path.exists():
        return _err(f"file not found: {filepath}")
    if not path.is_file():
        return _err(f"not a file: {filepath}")
    size = path.stat().st_size
    if size > 50 * 1024 * 1024:
        return _err(f"file too large ({size // 1024 // 1024} MB); limit is 50 MB")
    try:
        mime_type, _ = mimetypes.guess_type(str(path))
        if not mime_type:
            mime_type = "application/octet-stream"
        raw = path.read_bytes()
        file_base64 = base64.b64encode(raw).decode()
        return _ok({
            "filename": path.name,
            "mime_type": mime_type,
            "file_base64": file_base64,
        })
    except Exception as e:
        return _err(f"send-file error: {e}")


# ── Delegate to Claude Code (claude -p, on subscription) ─────────────────────

def _resolve_perm(payload: dict) -> str:
    """Pick claude -p permission mode based on caller-supplied trust signal.

    payload.trusted=True  → bypassPermissions (full autonomy: edits + shell).
    payload.trusted=False → acceptEdits        (edits only; safer default).
    Legacy: payload.cautious=True forces non-trusted regardless of `trusted`.
    """
    trusted = bool(payload.get("trusted", False))
    if payload.get("cautious"):
        trusted = False
    return "bypassPermissions" if trusted else "acceptEdits"


def run_claude_code_stream(payload: dict, on_event=None, on_proc=None) -> dict:
    """Same as handle_run_claude_code but with `--output-format stream-json`
    so we can surface progress to Telegram chunk-by-chunk.

    `on_event(evt)` is called with each parsed JSON event from stdout.
    `on_proc(p)`    is called once with the subprocess.Popen instance right
                    after spawn, so the caller can hand the user a kill button.
    Returns the same shape as handle_run_claude_code: {"success", "data": {result, cost_usd, model, num_turns}, ...}.
    """
    prompt = (payload.get("prompt") or "").strip()
    if not prompt:
        return _err("no prompt provided")
    model = (payload.get("model") or "sonnet").strip()
    cwd = payload.get("cwd") or None
    if cwd and not Path(cwd).is_dir():
        return _err(f"cwd not found: {cwd}")
    perm = _resolve_perm(payload)
    config_dir = payload.get("config_dir") or None
    timeout = int(payload.get("timeout", 1800))

    claude = shutil.which("claude")
    if not claude:
        return _err("claude CLI not found on PATH")

    env = dict(os.environ)
    env.pop("ANTHROPIC_API_KEY", None)
    if config_dir:
        env["CLAUDE_CONFIG_DIR"] = config_dir

    # stream-json + --verbose: NDJSON on stdout, one event per line.
    args = [claude, "-p", prompt, "--model", model,
            "--permission-mode", perm,
            "--output-format", "stream-json", "--verbose"]
    try:
        proc = subprocess.Popen(
            args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            env=env, cwd=cwd, text=True, encoding="utf-8", errors="replace",
            bufsize=1,  # line-buffered
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        )
    except Exception as e:
        return _err(f"claude -p не запустился: {e}")

    if on_proc:
        try:
            on_proc(proc)
        except Exception as e:
            # Caller's bookkeeping shouldn't break the run — just log.
            import logging
            logging.getLogger(__name__).warning(f"on_proc callback failed: {e}")

    # Drain stderr in a daemon thread so the OS pipe never fills and blocks
    # the child on stderr write (which would also stall our stdout readline).
    stderr_chunks: list[str] = []
    def _drain_stderr():
        try:
            for ln in iter(proc.stderr.readline, ""):
                stderr_chunks.append(ln)
        except Exception:
            pass
    threading.Thread(target=_drain_stderr, daemon=True, name="claude-stderr").start()

    # Hard wall-clock kill: fires even when readline() is blocked on a silent
    # hang (the in-loop timeout check only runs after each new stdout line).
    timed_out = {"v": False}
    def _hard_kill():
        timed_out["v"] = True
        try:
            proc.kill()
        except Exception:
            pass
    watchdog = threading.Timer(timeout, _hard_kill)
    watchdog.daemon = True
    watchdog.start()

    final = {"result": "", "cost_usd": None, "model": model, "num_turns": None}
    started = time.time()
    try:
        for line in iter(proc.stdout.readline, ""):
            line = line.strip()
            if not line:
                continue
            try:
                evt = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(evt, dict):
                continue
            if on_event:
                try:
                    on_event(evt)
                except Exception:
                    pass  # don't let UI errors kill the run
            if evt.get("type") == "result":
                final["result"]   = evt.get("result", "") or ""
                final["cost_usd"] = evt.get("total_cost_usd")
                final["num_turns"] = evt.get("num_turns")
                if evt.get("is_error") or evt.get("subtype") not in (None, "success"):
                    watchdog.cancel()
                    return _err(f"claude -p: {final['result'] or evt.get('subtype')}")
        proc.wait(timeout=10)
    except subprocess.TimeoutExpired:
        proc.kill()
        watchdog.cancel()
        return _err("claude -p не завершился после закрытия stdout")
    except Exception as e:
        proc.kill()
        watchdog.cancel()
        return _err(f"claude -p stream error: {e}")
    finally:
        watchdog.cancel()

    if timed_out["v"]:
        return _err(f"claude -p превысил таймаут {timeout}s")
    if proc.returncode != 0 and not final["result"]:
        err = "".join(stderr_chunks)
        return _err(f"claude -p rc={proc.returncode}: {err[:500]}")
    return _ok(final)


def handle_run_claude_code(payload: dict) -> dict:
    """Delegate a heavy task to Claude Code via `claude -p`.

    Runs on the user's SUBSCRIPTION (OAuth), not the paid per-token API:
    ANTHROPIC_API_KEY is stripped from the child env so Claude Code falls back
    to its logged-in account. `config_dir` selects which logged-in account to
    use via CLAUDE_CONFIG_DIR (for multi-account switching).
    """
    prompt = (payload.get("prompt") or "").strip()
    if not prompt:
        return _err("no prompt provided")

    model = (payload.get("model") or "sonnet").strip()
    cwd = payload.get("cwd") or None
    if cwd and not Path(cwd).is_dir():
        return _err(f"cwd not found: {cwd}")
    # Permission resolution: trusted cwds (allowlisted, or interactively picked
    # via /claude) get bypassPermissions for full autonomy; everything else
    # falls back to acceptEdits — claude -p can still edit files but no auto-shell.
    perm = _resolve_perm(payload)
    config_dir = payload.get("config_dir") or None
    timeout = int(payload.get("timeout", 900))

    claude = shutil.which("claude")
    if not claude:
        return _err("claude CLI not found on PATH")

    env = dict(os.environ)
    env.pop("ANTHROPIC_API_KEY", None)  # force subscription auth, not paid API
    if config_dir:
        env["CLAUDE_CONFIG_DIR"] = config_dir

    args = [claude, "-p", prompt, "--model", model,
            "--permission-mode", perm, "--output-format", "json"]
    try:
        r = subprocess.run(
            args, capture_output=True, text=True, timeout=timeout, env=env, cwd=cwd,
            encoding="utf-8", errors="replace",
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        )
    except subprocess.TimeoutExpired:
        return _err(f"claude -p превысил таймаут {timeout}s")
    except Exception as e:
        return _err(f"claude -p не запустился: {e}")

    out = (r.stdout or "").strip()
    try:
        parsed = json.loads(out)
    except json.JSONDecodeError:
        if r.returncode != 0:
            return _err(f"claude -p rc={r.returncode}: {(r.stderr or out)[:500]}")
        return _ok({"result": out[:6000], "cost_usd": None, "model": model})

    if parsed.get("is_error") or parsed.get("subtype") not in (None, "success"):
        return _err(f"claude -p: {parsed.get('result') or parsed.get('subtype')}")
    return _ok({
        "result": parsed.get("result", "") or "",
        "cost_usd": parsed.get("total_cost_usd"),
        "model": model,
        "num_turns": parsed.get("num_turns"),
    })


# ── Read / extract text from any document format ─────────────────────────────

_PLAIN_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".log", ".json", ".yaml", ".yml",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".sql", ".sh", ".ps1", ".bat", ".ini",
    ".cfg", ".conf", ".toml", ".env", ".srt", ".vtt",
}


def _read_text_file(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


def _looks_binary(path: Path, sniff: int = 8192) -> bool:
    """Sniff raw bytes to decide if a file is binary. Done on bytes (not via the
    all-accepting latin-1 decode) so binaries are routed to the claude -p fallback
    instead of returning mojibake."""
    chunk = path.read_bytes()[:sniff]
    if not chunk:
        return False  # empty -> treat as (trivial) text
    if b"\x00" in chunk:
        return True
    # text in any encoding has few control chars; lots of them => binary
    control = sum(1 for b in chunk if b < 0x09 or 0x0e <= b <= 0x1f)
    return control / len(chunk) > 0.30


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n".join((pg.extract_text() or "") for pg in reader.pages)


def _extract_docx(path: Path) -> str:
    import docx
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _extract_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            out.append(f"# Лист: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    out.append("\t".join(cells))
    finally:
        wb.close()
    return "\n".join(out)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"# Слайд {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(r.text for r in para.runs)
                    if txt.strip():
                        out.append(txt)
            if shape.has_table:
                for row in shape.table.rows:
                    out.append("\t".join(c.text for c in row.cells))
    return "\n".join(out)


def _extract_html(path: Path) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(_read_text_file(path), "html.parser")
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text("\n", strip=True)


def _extract_xml(path: Path) -> str:
    raw = _read_text_file(path)
    try:
        import xml.etree.ElementTree as ET
        root = ET.fromstring(raw)
        parts = []
        for el in root.iter():
            txt = (el.text or "").strip()
            if txt:
                tag = el.tag.split("}")[-1]  # strip namespace
                parts.append(f"{tag}: {txt}")
        return "\n".join(parts) or raw
    except Exception:
        return raw


def _extract_rtf(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text
    return rtf_to_text(_read_text_file(path))


def _read_via_claude(path: Path) -> dict:
    """Fallback for formats with no native extractor (.doc/.xls/.ppt/odf/etc.):
    delegate reading to Claude Code, which can shell out to convert them."""
    res = handle_run_claude_code({
        "prompt": (f"Прочитай файл {path} и верни ТОЛЬКО его текстовое содержимое, "
                   "без своих комментариев. Если это таблица/презентация/документ — "
                   "извлеки весь читаемый текст."),
        "cwd": str(path.parent),
        "model": "sonnet",
        "cautious": True,
    })
    if res.get("success"):
        text = res["data"].get("result", "")
        return _ok({"filename": path.name, "ext": path.suffix.lower(),
                    "chars": len(text), "truncated": False, "text": text,
                    "via": "claude-code"})
    return res


def handle_read_document(payload: dict) -> dict:
    """Extract text from virtually any document format.

    Native extractors for pdf/docx/xlsx/pptx/html/xml/rtf + all plain-text types;
    legacy/binary office (.doc/.xls/.ppt) and unknown binaries fall back to
    Claude Code. Returns text (truncated to max_chars for chat)."""
    # Same alias-tolerance as send-file — Haiku occasionally swaps in `path` etc.
    fp = (payload.get("filepath") or payload.get("path")
          or payload.get("file") or payload.get("file_path") or "").strip()
    if not fp:
        return _err("no filepath provided")
    path = Path(fp)
    if not path.exists():
        return _err(f"file not found: {fp}")
    if not path.is_file():
        return _err(f"not a file: {fp}")

    max_chars = int(payload.get("max_chars", 20000))
    ext = path.suffix.lower()
    fallback_exts = {".doc", ".xls", ".ppt", ".odt", ".ods", ".odp",
                     ".pages", ".key", ".numbers", ".epub"}

    try:
        if ext == ".pdf":
            text = _extract_pdf(path)
        elif ext == ".docx":
            text = _extract_docx(path)
        elif ext in (".xlsx", ".xlsm"):
            text = _extract_xlsx(path)
        elif ext == ".pptx":
            text = _extract_pptx(path)
        elif ext in (".html", ".htm"):
            text = _extract_html(path)
        elif ext == ".xml":
            text = _extract_xml(path)
        elif ext == ".rtf":
            text = _extract_rtf(path)
        elif ext in _PLAIN_TEXT_EXTS:
            text = _read_text_file(path)
        elif ext in fallback_exts:
            return _read_via_claude(path)
        else:
            # unknown / no extension: sniff raw bytes first; delegate if binary
            if _looks_binary(path):
                return _read_via_claude(path)
            text = _read_text_file(path)
    except Exception as e:
        fb = _read_via_claude(path)
        if fb.get("success"):
            return fb
        return _err(f"не удалось извлечь текст ({ext}): {e}")

    text = text or ""
    return _ok({"filename": path.name, "ext": ext, "chars": len(text),
                "truncated": len(text) > max_chars, "text": text[:max_chars]})


# ── Dispatcher ────────────────────────────────────────────────────────────────

def dispatch(command: dict, apps: dict, obsidian_cfg: dict | None = None) -> dict:
    cmd_type = command.get("type", "")
    payload  = command.get("payload", {})
    obs_cfg  = obsidian_cfg or {}

    if cmd_type == "screenshot":
        return handle_screenshot(payload)
    elif cmd_type == "computer":
        return handle_computer(payload)
    elif cmd_type == "terminal":
        return handle_terminal(payload)
    elif cmd_type == "launch-app":
        return handle_launch_app(payload, apps)
    elif cmd_type == "chrome-profile":
        return handle_chrome_profile(payload, apps)
    elif cmd_type == "chrome-profiles-list":
        return handle_chrome_profiles_list(payload)
    elif cmd_type == "install":
        return handle_install(payload)
    elif cmd_type == "system-info":
        return handle_system_info(payload)
    elif cmd_type == "obsidian-log":
        return handle_obsidian_log(payload, obs_cfg)
    elif cmd_type == "obsidian-read":
        return handle_obsidian_read(payload, obs_cfg)
    elif cmd_type == "obsidian-context":
        return handle_obsidian_context(payload, obs_cfg)
    elif cmd_type == "send-file":
        return handle_send_file(payload)
    elif cmd_type == "run-claude-code":
        return handle_run_claude_code(payload)
    elif cmd_type == "read-document":
        return handle_read_document(payload)
    else:
        return _err(f"unknown command type: {cmd_type}")
